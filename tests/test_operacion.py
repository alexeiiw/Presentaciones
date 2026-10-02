from __future__ import annotations

import os
import subprocess
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from PIL import Image

from motor import biblioteca_imagenes, proveedores_imagenes
from motor.estilos import nombres_estilos
from motor.exportadores import convertir_pptx_a_pdf
from motor.generador_clases import contar_diapositivas_estimadas, generar_presentacion
from motor.parser_markdown import parsear_markdown
from motor.validaciones import validar_clase
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]


class RespuestaFalsa:
    def __init__(self, data=None, content=b"imagen"):
        self._data = data or {}
        self.content = content

    def raise_for_status(self):
        return None

    def json(self):
        return self._data


class PruebasDeOperacion(unittest.TestCase):
    @contextmanager
    def _directorio_temporal(self, ruta: Path):
        anterior = Path.cwd()
        os.chdir(ruta)
        try:
            yield
        finally:
            os.chdir(anterior)

    def test_ejemplo_de_repositorio_pasa_revision_sin_errores(self):
        archivo = ROOT / "ejemplos" / "redes_computadoras.md"
        markdown = archivo.read_text(encoding="utf-8")
        clase = parsear_markdown(markdown)
        resultado = validar_clase(markdown, clase, ROOT)

        self.assertTrue(resultado.valido, resultado.errores)
        self.assertGreaterEqual(resultado.diapositivas_estimadas, len(clase.diapositivas))
        self.assertFalse(any("no parece tener una diapositiva desarrollada" in aviso for aviso in resultado.advertencias))
        self.assertFalse(any("no parece estar desarrollado" in aviso for aviso in resultado.advertencias))

    def test_validacion_distingue_errores_y_recomendaciones(self):
        markdown = """# Clase: Control de validación

Contenido Presentacion:
- Redes y seguridad

## Diapositiva: Inicio

Tipo: valor_inventado
"""
        resultado = validar_clase(markdown)
        self.assertFalse(resultado.valido)
        self.assertTrue(any("no soportado" in error for error in resultado.errores))
        self.assertTrue(any("no parece tener" in aviso for aviso in resultado.advertencias))

    def test_detecta_secciones_sin_desarrollo_y_bloques_sin_cierre(self):
        markdown = """# Clase: Secciones

## Diapositiva: Cambio de tema
Tipo: seccion

## Diapositiva: Segunda sección
Tipo: seccion

## Diapositiva: Código
Tipo: codigo
```python
print('hola')
"""
        resultado = validar_clase(markdown)
        self.assertTrue(any("no va seguida" in aviso for aviso in resultado.advertencias))
        self.assertTrue(any("sin cierre" in aviso for aviso in resultado.advertencias))
        self.assertTrue(resultado.valido)

    def test_parser_acepta_fechas_pendientes_y_codigo_sin_cierre(self):
        markdown = """# Clase: Parser

Fechas de Parciales:
- Parcial final: 30 de octubre

Qué queda pendiente:
- Revisar direccionamiento IPv6

## Diapositiva: Código
```python
print('fin')
"""
        clase = parsear_markdown(markdown)
        self.assertEqual(clase.fechas_parciales, ["Parcial final: 30 de octubre"])
        self.assertEqual(clase.pendientes, ["Revisar direccionamiento IPv6"])
        self.assertEqual(clase.diapositivas[0].codigo, "print('fin')")

    def test_generacion_cubre_layouts_y_cierre_en_todos_los_estilos(self):
        markdown = """# Clase: Suite de layouts
Agenda:
- Contexto
Contenido Presentacion:
- Componentes de red
Fechas Parciales:
- Parcial 1: 15 de octubre
Aprendizajes:
- Explicar componentes de red.
Qué queda pendiente:
- Estudiar enrutamiento avanzado.

## Diapositiva: Contenido general
Objetivo: Comprender la arquitectura.
Contenido:
- Un router conecta redes.
- Un switch conecta equipos.

## Diapositiva: Comparación
Tipo: columnas
Contenido:
- TCP prioriza fiabilidad.
- UDP prioriza baja latencia.

## Diapositiva: Camino de un paquete
Tipo: ruta
Contenido:
- Aplicación
- Transporte
- Red

## Diapositiva: Principio de red
Tipo: frase
Objetivo: Una red conecta equipos para intercambiar información.

## Diapositiva: Arquitectura
Tipo: diagrama
Diagrama: bloques
Contenido:
- Cliente: origina la solicitud.
- Router: conecta las redes.

## Diapositiva: Código
Tipo: codigo
```python
print('red')
```

## Diapositiva: Material de consulta
Tipo: repositorio
Contenido:
- Documentación técnica oficial.

## Diapositiva: Práctica previa heredada
Tipo: actividad
Contenido:
- Analizar una topología.
"""
        clase = parsear_markdown(markdown)
        with tempfile.TemporaryDirectory() as temporal:
            for estilo in nombres_estilos():
                salida = Path(temporal) / f"{estilo}.pptx"
                generar_presentacion(clase, salida, estilo_nombre=estilo, modo_imagen="sin_imagen")
                self.assertTrue(salida.exists(), estilo)
                presentacion = Presentation(salida)
                # Portada + agenda + índice + 8 diapositivas fuente + parcial + aprendizajes + pendientes.
                self.assertEqual(len(presentacion.slides), 14, estilo)
                titulos = [shape.text for slide in presentacion.slides for shape in slide.shapes if shape.has_text_frame]
                self.assertTrue(any("Suite de layouts" in titulo for titulo in titulos), estilo)
                self.assertTrue(any("Fechas de parciales" in titulo for titulo in titulos), estilo)
                self.assertTrue(any("Qué aprendiste" in titulo for titulo in titulos), estilo)
                self.assertTrue(any("Qué queda pendiente" in titulo for titulo in titulos), estilo)

    def test_generacion_funciona_desde_otro_directorio_y_resuelve_assets(self):
        markdown = """# Clase: Rutas independientes
Imagen Portada: images.jpg
Logo Portada: Umg.png

## Diapositiva: Prueba local
Tipo: contenido
Contenido:
- Las rutas resuelven desde la raíz del proyecto.
"""
        clase = parsear_markdown(markdown)
        with tempfile.TemporaryDirectory() as temporal:
            salida = Path(temporal) / "fuera_del_proyecto.pptx"
            with patch("pathlib.Path.cwd", return_value=Path(temporal)):
                generar_presentacion(clase, salida, modo_imagen="sin_imagen")
            presentacion = Presentation(salida)
            imagenes_incrustadas = [shape for slide in presentacion.slides for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(len(imagenes_incrustadas), 2)

    def test_mismo_nombre_reemplaza_pptx_y_actualiza_contenido(self):
        salida = Path(tempfile.gettempdir()) / "presentaciones_test_reemplazo.pptx"
        try:
            for titulo in ("Primera versión", "Segunda versión"):
                clase = parsear_markdown(f"# Clase: {titulo}\n\n## Diapositiva: Tema\n\nContenido:\n- Punto\n")
                generar_presentacion(clase, salida, modo_imagen="sin_imagen")
                presentacion = Presentation(salida)
                textos = [shape.text for shape in presentacion.slides[0].shapes if shape.has_text_frame]
                self.assertIn(titulo, textos)
        finally:
            salida.unlink(missing_ok=True)

    def test_proveedores_no_reusan_resultados_agotados(self):
        usadas = set()
        fotos = [{"id": 101}, {"id": 202}]
        primera = proveedores_imagenes._elegir_no_usada(fotos, usadas, "pexels")
        segunda = proveedores_imagenes._elegir_no_usada(fotos, usadas, "pexels")
        tercera = proveedores_imagenes._elegir_no_usada(fotos, usadas, "pexels")
        self.assertIsNotNone(primera)
        self.assertIsNotNone(segunda)
        self.assertNotEqual(primera["id"], segunda["id"])
        self.assertIsNone(tercera)

    def test_busqueda_de_proveedor_entrega_imagen_unica(self):
        destino = Path(tempfile.gettempdir()) / "prueba_proveedor_presentaciones"
        usadas = set()
        respuesta_api = RespuestaFalsa({"photos": [
            {"id": 7, "src": {"large2x": "https://img/7.jpg"}},
            {"id": 8, "src": {"large2x": "https://img/8.jpg"}},
        ]})
        respuesta_img = RespuestaFalsa(content=b"fake-jpeg")
        try:
            with patch.dict("os.environ", {"PEXELS_API_KEY": "prueba"}), patch.object(
                proveedores_imagenes.requests, "get", side_effect=[respuesta_api, respuesta_img, respuesta_api, respuesta_img]
            ):
                primera = proveedores_imagenes.buscar_imagen_pexels("network router", destino, usadas)
                segunda = proveedores_imagenes.buscar_imagen_pexels("network router", destino, usadas)
            self.assertIsNotNone(primera)
            self.assertIsNotNone(segunda)
            self.assertNotEqual(primera[2], segunda[2])
            self.assertEqual(len(list(destino.glob("*.jpg"))), 2)
        finally:
            for imagen in destino.glob("*.jpg"):
                imagen.unlink()
            destino.rmdir() if destino.exists() else None

    def test_biblioteca_usa_rutas_estables_y_no_pierde_imagenes(self):
        with tempfile.TemporaryDirectory() as temporal:
            raiz = Path(temporal)
            origen = raiz / "origen.jpg"
            origen.write_bytes(b"test-image")
            with patch.multiple(
                biblioteca_imagenes,
                BASE_BIBLIOTECA=raiz / "biblioteca_imagenes",
                DIR_IMAGENES=raiz / "biblioteca_imagenes" / "imagenes",
                INDEX_PATH=raiz / "biblioteca_imagenes" / "index.json",
            ):
                imagen = biblioteca_imagenes.guardar_en_biblioteca("redes y seguridad", origen, "local")
                registro = biblioteca_imagenes.leer_index()[0]
                self.assertEqual(Path(registro["archivo"]).parts[0], "biblioteca_imagenes")
                self.assertTrue(imagen.exists())
                (raiz / "otro_directorio").mkdir()
                with self._directorio_temporal(raiz / "otro_directorio"):
                    seleccionada = biblioteca_imagenes.buscar_en_biblioteca("redes y seguridad")
                self.assertEqual(seleccionada, imagen.resolve())
                self.assertEqual(biblioteca_imagenes.buscar_en_biblioteca("redes y seguridad"), imagen.resolve())
                self.assertEqual(biblioteca_imagenes._normalizar_keyword("ciberseguridad y redes"), "ciberseguridad_y_redes")

    def test_actualiza_estados_masivos_sin_omitir_registros(self):
        with tempfile.TemporaryDirectory() as temporal:
            raiz = Path(temporal)
            indice = raiz / "biblioteca_imagenes" / "index.json"
            with patch.multiple(
                biblioteca_imagenes,
                BASE_BIBLIOTECA=raiz / "biblioteca_imagenes",
                DIR_IMAGENES=raiz / "biblioteca_imagenes" / "imagenes",
                INDEX_PATH=indice,
            ):
                (raiz / "biblioteca_imagenes" / "imagenes").mkdir(parents=True)
                (raiz / "biblioteca_imagenes" / "imagenes" / "uno.jpg").write_bytes(b"uno")
                (raiz / "biblioteca_imagenes" / "imagenes" / "dos.jpg").write_bytes(b"dos")
                biblioteca_imagenes.guardar_index(
                    [
                        {"archivo": "biblioteca_imagenes/imagenes/uno.jpg", "estado": "pendiente"},
                        {"archivo": "biblioteca_imagenes/imagenes/dos.jpg", "estado": "pendiente"},
                        {"archivo": "biblioteca_imagenes/imagenes/tres.jpg", "estado": "aprobada"},
                    ]
                )
                uno = raiz / "biblioteca_imagenes" / "imagenes" / "uno.jpg"
                dos = raiz / "biblioteca_imagenes" / "imagenes" / "dos.jpg"
                total = biblioteca_imagenes.actualizar_estados_archivos([str(uno), str(dos), str(uno)], "rechazada")
                estados = [registro["estado"] for registro in biblioteca_imagenes.leer_index()]
                self.assertEqual(total, 2)
                self.assertEqual(estados, ["rechazada", "rechazada", "aprobada"])
                with self.assertRaises(ValueError):
                    biblioteca_imagenes.actualizar_estados_archivos([str(uno)], "publicada")

    def test_generador_de_miniaturas_produce_uri_pequeno_para_galeria(self):
        with tempfile.TemporaryDirectory() as temporal:
            imagen_path = Path(temporal) / "imagen.png"
            Image.new("RGB", (1600, 1200), color=(30, 90, 150)).save(imagen_path)
            from app import imagen_galeria_data_uri

            uri = imagen_galeria_data_uri(imagen_path)
            self.assertTrue(uri.startswith("data:image/jpeg;base64,"))
            self.assertLess(len(uri), 12000)

    def test_interfaz_abre_valida_y_genera_con_el_flujo_normal(self):
        nombre_prueba = "prueba_interfaz_automatizada"
        salida = ROOT / "salidas" / f"{nombre_prueba}.pptx"
        salida_pdf = salida.with_suffix(".pdf")
        previo = salida.read_bytes() if salida.exists() else None
        previo_pdf = salida_pdf.read_bytes() if salida_pdf.exists() else None

        def ejecutar_libreoffice_falso(comando, **kwargs):
            carpeta_salida = Path(comando[comando.index("--outdir") + 1])
            (carpeta_salida / salida_pdf.name).write_bytes(b"pdf-generado")
            return subprocess.CompletedProcess(comando, 0, stdout="", stderr="")

        try:
            app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
            self.assertFalse(app.exception, [str(error.message) for error in app.exception])
            self.assertEqual(len(app.metric), 3)
            self.assertFalse(app.button[0].disabled)
            app.text_input[0].set_value(nombre_prueba).run()
            app.button[0].click().run()
            self.assertFalse(app.exception, [str(error.message) for error in app.exception])
            self.assertTrue(any("Presentacion generada" in aviso.value for aviso in app.success))
            self.assertTrue(salida.exists())
            self.assertTrue(Presentation(salida).slides)
            self.assertEqual(len(app.get("download_button")), 1)
            self.assertTrue(any(boton.key == "convertir_pptx_actual" for boton in app.button))
            app.text_input[0].set_value("otro_nombre").run()
            self.assertTrue(any("configuración cambió" in aviso.value for aviso in app.info))
            self.assertFalse(any(boton.label.startswith("Descargar") for boton in app.button))
            app.text_input[0].set_value(nombre_prueba).run()
            self.assertTrue(any(boton.key == "convertir_pptx_actual" for boton in app.button))
            with patch("motor.exportadores.detectar_libreoffice", return_value="soffice"), patch(
                "motor.exportadores.subprocess.run", side_effect=ejecutar_libreoffice_falso
            ):
                boton_pdf = next(boton for boton in app.button if boton.key == "convertir_pptx_actual")
                boton_pdf.click().run()
                self.assertFalse(app.exception, [str(error.message) for error in app.exception])
                self.assertTrue(salida_pdf.exists())
                self.assertGreaterEqual(len(app.get("download_button")), 2)
            selector_estilo = next(
                selector for selector in app.selectbox if selector.label == "Estilo visual"
            )
            selector_estilo.set_value("Academico Formal").run()
            self.assertFalse(any(boton.label == "Descargar PPTX" for boton in app.button))
        finally:
            if previo is None:
                salida.unlink(missing_ok=True)
            else:
                salida.write_bytes(previo)
            if previo_pdf is None:
                salida_pdf.unlink(missing_ok=True)
            else:
                salida_pdf.write_bytes(previo_pdf)

    def test_curador_muestra_galeria_seleccion_multiple_y_acciones_lote(self):
        with tempfile.TemporaryDirectory() as temporal:
            raiz = Path(temporal)
            biblioteca = raiz / "biblioteca_imagenes"
            imagenes = biblioteca / "imagenes"
            imagenes.mkdir(parents=True)
            imagen_a = imagenes / "redes_0001.jpg"
            imagen_b = imagenes / "seguridad_0002.jpg"
            Image.new("RGB", (40, 30), color=(20, 80, 120)).save(imagen_a)
            Image.new("RGB", (40, 30), color=(120, 40, 80)).save(imagen_b)
            archivo_indice = biblioteca / "index.json"
            original_base = biblioteca_imagenes.BASE_BIBLIOTECA
            original_imagenes = biblioteca_imagenes.DIR_IMAGENES
            original_index = biblioteca_imagenes.INDEX_PATH
            try:
                with patch.multiple(
                    biblioteca_imagenes,
                    BASE_BIBLIOTECA=biblioteca,
                    DIR_IMAGENES=imagenes,
                    INDEX_PATH=archivo_indice,
                ):
                    biblioteca_imagenes.guardar_index(
                        [
                            {"keyword": "network router", "keyword_normalizada": "network_router", "archivo": str(imagen_a), "proveedor": "pexels", "url": "https://example.test/a", "usos": 1, "estado": "pendiente"},
                            {"keyword": "cyber security", "keyword_normalizada": "cyber_security", "archivo": str(imagen_b), "proveedor": "pixabay", "url": "https://example.test/b", "usos": 2, "estado": "pendiente"},
                        ]
                    )
                    with patch("app.leer_index", side_effect=biblioteca_imagenes.leer_index):
                        app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
                        filtro = next(radio for radio in app.radio if radio.label == "Filtrar imágenes por estado")
                        self.assertIn("todas", filtro.options)
                        seleccionar = next(
                            boton for boton in app.button if boton.key == "seleccionar_todas_imagenes"
                        )
                        seleccionar.click().run()
                        aprobar = next(boton for boton in app.button if boton.key == "aprobar_imagenes_masivo")
                        self.assertFalse(aprobar.disabled)
                        aprobar.click().run()
                        self.assertFalse(app.exception, [str(error.message) for error in app.exception])
                        estados = [registro["estado"] for registro in biblioteca_imagenes.leer_index()]
                        self.assertEqual(estados, ["aprobada", "aprobada"])
            finally:
                biblioteca_imagenes.BASE_BIBLIOTECA = original_base
                biblioteca_imagenes.DIR_IMAGENES = original_imagenes
                biblioteca_imagenes.INDEX_PATH = original_index

    def test_interfaz_muestra_pptx_de_sesion_anterior_y_permite_convertirlo(self):
        nombre = "presentacion_sesion_anterior_test"
        pptx = ROOT / "salidas" / f"{nombre}.pptx"
        pdf = pptx.with_suffix(".pdf")
        pptx_anterior = pptx.read_bytes() if pptx.exists() else None
        pdf_anterior = pdf.read_bytes() if pdf.exists() else None

        def ejecutar_libreoffice_falso(comando, **kwargs):
            directorio_salida = Path(comando[comando.index("--outdir") + 1])
            (directorio_salida / f"{pptx.stem}.pdf").write_bytes(b"pdf-convertido")
            return subprocess.CompletedProcess(comando, 0, stdout="", stderr="")

        try:
            Presentation().save(pptx)
            with patch("motor.exportadores.detectar_libreoffice", return_value="soffice"), patch(
                "motor.exportadores.subprocess.run", side_effect=ejecutar_libreoffice_falso
            ) as conversor:
                app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
                selector = next(select for select in app.selectbox if select.label == "Presentación guardada")
                etiqueta = next(opcion for opcion in selector.options if nombre in opcion)
                self.assertIn(nombre, etiqueta)
                self.assertTrue(any("Convertir PPTX a PDF" in tab.label for tab in app.tabs))
                boton = next(boton for boton in app.button if boton.key == "convertir_pptx_guardado")
                boton.click().run()
                self.assertFalse(app.exception, [str(error.message) for error in app.exception])
                conversor.assert_called_once()
                self.assertTrue(pdf.exists())
                self.assertEqual(len(app.get("download_button")), 1)
                self.assertTrue(any("PDF generado" in exito.value for exito in app.success))
        finally:
            if pptx_anterior is None:
                pptx.unlink(missing_ok=True)
            else:
                pptx.write_bytes(pptx_anterior)
            if pdf_anterior is None:
                pdf.unlink(missing_ok=True)
            else:
                pdf.write_bytes(pdf_anterior)

    def test_convertidor_permite_regenerar_pdf_de_pptx_existente(self):
        with tempfile.TemporaryDirectory() as temporal:
            pptx = Path(temporal) / "clase.pptx"
            pdf = Path(temporal) / "clase.pdf"
            Presentation().save(pptx)
            pdf.write_bytes(b"pdf-anterior")

            def convertir_falso(comando, **kwargs):
                temporal = Path(comando[comando.index("--outdir") + 1]) / "clase.pdf"
                temporal.write_bytes(b"pdf-actualizado")
                return subprocess.CompletedProcess(comando, 0, stdout="", stderr="")

            with patch("motor.exportadores.detectar_libreoffice", return_value="soffice"), patch(
                "motor.exportadores.subprocess.run", side_effect=convertir_falso
            ):
                resultado = convertir_pptx_a_pdf(pptx, pdf)

            self.assertTrue(resultado.ok, resultado.mensaje)
            self.assertEqual(resultado.ruta, pdf)
            self.assertEqual(pdf.read_bytes(), b"pdf-actualizado")

    def test_claves_locales_no_se_incluyen_en_el_diff_publicable(self):
        resultado = subprocess.run(
            ["git", "status", "--short", "--untracked-files=all"],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        archivos_cambiados = {linea[3:] for linea in resultado.stdout.splitlines() if len(linea) > 3}
        self.assertFalse(any("claves.txt" in archivo or archivo.endswith(".env") for archivo in archivos_cambiados))


if __name__ == "__main__":
    unittest.main()
